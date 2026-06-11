from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MD_PATH = next((ROOT / "docs").glob("*.md"))
OUT_PATH = ROOT / "docs" / f"{MD_PATH.stem}.docx"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(85, 85, 85)
RED = RGBColor(179, 32, 32)
LIGHT_GRAY = "F2F4F7"
BORDER = "DADCE0"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_paragraph(paragraph, before=0, after=6, line=1.10, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_para(
    doc,
    text="",
    style=None,
    size=11,
    color=None,
    bold=False,
    italic=False,
    after=6,
    before=0,
    line=1.10,
    align=None,
):
    paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    configure_paragraph(paragraph, before=before, after=after, line=line, align=align)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10


def add_header_footer(doc):
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.text = ""
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    configure_paragraph(header_p, after=0, line=1.0)
    set_run_font(header_p.add_run("首经贸留学助手 | 作品设计书"), size=9, color=GRAY)

    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_paragraph(footer_p, after=0, line=1.0)
    set_run_font(
        footer_p.add_run("Vercel AI 版 + GitHub Pages 无 AI 版 | React + TypeScript + Tailwind CSS"),
        size=8.5,
        color=GRAY,
    )


def add_metadata_table(doc):
    rows = [
        ("作品名称", "首经贸留学助手"),
        ("作品类型", "网页应用 / 留学选校分析系统"),
        ("目标用户", "首都经济贸易大学普通本科生"),
        ("技术栈", "React + TypeScript + Tailwind CSS + Vite"),
        ("AI 版本", "Vercel Serverless API + DeepSeek Chat"),
        ("静态版本", "GitHub Pages 本地规则分析报告"),
        ("项目地址", "https://cueb-study-abroad-assistant.vercel.app/"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    set_table_borders(table)
    for row_index, (key, value) in enumerate(rows):
        key_cell, value_cell = table.rows[row_index].cells
        for cell in (key_cell, value_cell):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        set_cell_shading(key_cell, LIGHT_GRAY)
        key_cell.width = Inches(1.35)
        value_cell.width = Inches(5.15)

        key_cell.text = ""
        key_p = key_cell.paragraphs[0]
        configure_paragraph(key_p, after=0, line=1.1)
        set_run_font(key_p.add_run(key), size=10.5, color=NAVY, bold=True)

        value_cell.text = ""
        value_p = value_cell.paragraphs[0]
        configure_paragraph(value_p, after=0, line=1.1)
        set_run_font(value_p.add_run(value), size=10.5, color=RGBColor(25, 25, 25))


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table)
    set_table_borders(table, color="E0E7EF")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    title_p = cell.paragraphs[0]
    configure_paragraph(title_p, after=4, line=1.15)
    set_run_font(title_p.add_run(title), size=11.5, color=NAVY, bold=True)
    body_p = cell.add_paragraph()
    configure_paragraph(body_p, after=0, line=1.15)
    set_run_font(body_p.add_run(body), size=10.5, color=RGBColor(45, 55, 72))
    add_para(doc, "", after=4)


def add_cover(doc):
    add_para(doc, "作品设计书", size=11, color=RED, bold=True, after=4)
    add_para(doc, "首经贸留学助手", size=25, color=NAVY, bold=True, after=6)
    add_para(doc, "AI 驱动，基于真实案例的首经贸专属留学选校分析系统", size=13.5, color=GRAY, after=18)
    add_metadata_table(doc)
    add_para(doc, "", after=8)
    add_callout(
        doc,
        "设计摘要",
        "本作品面向首经贸普通本科生，围绕海外硕士申请中的背景定位、案例参考、选校分层、学校官网核对与时间规划，构建一套可解释的留学选校分析系统。系统保留 Vercel AI 版与无 AI 静态版两种交付形态，兼顾功能完整性与可访问性。",
    )
    doc.add_page_break()


def add_code_block(doc, lines):
    for line in lines:
        paragraph = doc.add_paragraph()
        configure_paragraph(paragraph, before=0, after=2, line=1.0)
        run = paragraph.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=9.2, color=RGBColor(40, 40, 40))
        paragraph.paragraph_format.left_indent = Inches(0.2)


def clean_inline(text):
    return text.replace("`", "")


def add_markdown_body(doc, text):
    in_code = False
    code_lines = []
    skip_first_h1 = True
    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped:
            continue
        if stripped.startswith("# "):
            if skip_first_h1:
                skip_first_h1 = False
                continue
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.add_run(clean_inline(stripped[2:]))
            for run in paragraph.runs:
                set_run_font(run, size=16, color=BLUE, bold=True)
        elif stripped.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.add_run(clean_inline(stripped[3:]))
            for run in paragraph.runs:
                set_run_font(run, size=16, color=BLUE, bold=True)
        elif stripped.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.add_run(clean_inline(stripped[4:]))
            for run in paragraph.runs:
                set_run_font(run, size=13, color=BLUE, bold=True)
        elif stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            configure_paragraph(paragraph, after=4, line=1.167)
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            run = paragraph.add_run(clean_inline(stripped[2:]))
            set_run_font(run, size=11)
        elif len(stripped) > 2 and stripped[0].isdigit() and ". " in stripped[:5]:
            paragraph = doc.add_paragraph(style="List Number")
            configure_paragraph(paragraph, after=4, line=1.167)
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            content = stripped.split(". ", 1)[1]
            run = paragraph.add_run(clean_inline(content))
            set_run_font(run, size=11)
        else:
            add_para(doc, clean_inline(stripped), size=11, after=6, line=1.10)


def main():
    text = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    style_doc(doc)
    add_header_footer(doc)
    add_cover(doc)
    add_markdown_body(doc, text)
    doc.save(OUT_PATH)
    print(str(OUT_PATH))


if __name__ == "__main__":
    main()
